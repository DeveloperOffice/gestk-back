import io
from datetime import datetime
from django.http import HttpResponse
from django.template.loader import render_to_string
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


class PDFExporter:
    """Classe para exportação de dados para PDF"""
    
    def __init__(self, title="Relatório GESTK"):
        self.title = title
        self.styles = getSampleStyleSheet()
        self._setup_custom_styles()
    
    def _setup_custom_styles(self):
        """Configura estilos personalizados para o PDF"""
        self.styles.add(ParagraphStyle(
            name='CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=18,
            spaceAfter=30,
            alignment=TA_CENTER,
            textColor=colors.darkblue
        ))
        
        self.styles.add(ParagraphStyle(
            name='CustomHeading',
            parent=self.styles['Heading2'],
            fontSize=14,
            spaceAfter=12,
            textColor=colors.darkblue
        ))
        
        self.styles.add(ParagraphStyle(
            name='CustomNormal',
            parent=self.styles['Normal'],
            fontSize=10,
            spaceAfter=6
        ))
    
    def export_carteira(self, carteira_data, contabilidade_nome):
        """Exporta dados da carteira para PDF"""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=18)
        
        story = []
        
        # Título
        story.append(Paragraph(f"Relatório de Carteira - {contabilidade_nome}", self.styles['CustomTitle']))
        story.append(Spacer(1, 20))
        
        # Data de geração
        story.append(Paragraph(f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}", self.styles['CustomNormal']))
        story.append(Spacer(1, 20))
        
        # Resumo geral
        if 'resumo' in carteira_data:
            story.append(Paragraph("Resumo Geral", self.styles['CustomHeading']))
            resumo = carteira_data['resumo']
            
            resumo_data = [
                ['Total de Empresas', str(resumo.get('total_empresas', 0))],
                ['Empresas Ativas', str(resumo.get('empresas_ativas', 0))],
                ['Empresas Inativas', str(resumo.get('empresas_inativas', 0))],
                ['Total de Funcionários', str(resumo.get('total_funcionarios', 0))],
                ['Faturamento Total', f"R$ {resumo.get('faturamento_total', 0):,.2f}"]
            ]
            
            resumo_table = Table(resumo_data, colWidths=[3*inch, 2*inch])
            resumo_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            story.append(resumo_table)
            story.append(Spacer(1, 20))
        
        # Lista de empresas
        if 'empresas' in carteira_data:
            story.append(Paragraph("Lista de Empresas", self.styles['CustomHeading']))
            
            empresas_data = [['Razão Social', 'CNPJ', 'Status', 'Funcionários']]
            for empresa in carteira_data['empresas'][:50]:  # Limitar a 50 empresas
                empresas_data.append([
                    empresa.get('razao_social', ''),
                    empresa.get('cnpj', ''),
                    'Ativa' if empresa.get('ativo', False) else 'Inativa',
                    str(empresa.get('total_funcionarios', 0))
                ])
            
            empresas_table = Table(empresas_data, colWidths=[2.5*inch, 1.5*inch, 1*inch, 1*inch])
            empresas_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            story.append(empresas_table)
        
        doc.build(story)
        buffer.seek(0)
        return buffer
    
    def export_clientes(self, clientes_data, contabilidade_nome):
        """Exporta dados de clientes para PDF"""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=18)
        
        story = []
        
        # Título
        story.append(Paragraph(f"Relatório de Clientes - {contabilidade_nome}", self.styles['CustomTitle']))
        story.append(Spacer(1, 20))
        
        # Data de geração
        story.append(Paragraph(f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}", self.styles['CustomNormal']))
        story.append(Spacer(1, 20))
        
        # Lista de clientes
        if clientes_data:
            story.append(Paragraph("Lista de Clientes", self.styles['CustomHeading']))
            
            clientes_table_data = [['Nome', 'Documento', 'Faturamento', 'Notas']]
            for cliente in clientes_data[:100]:  # Limitar a 100 clientes
                clientes_table_data.append([
                    cliente.get('nome', ''),
                    cliente.get('documento', ''),
                    f"R$ {cliente.get('total_transacoes', 0):,.2f}",
                    str(cliente.get('quantidade_transacoes', 0))
                ])
            
            clientes_table = Table(clientes_table_data, colWidths=[2.5*inch, 1.5*inch, 1.5*inch, 1*inch])
            clientes_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            story.append(clientes_table)
        
        doc.build(story)
        buffer.seek(0)
        return buffer
    
    def export_relatorio_geral(self, dados_gerais, contabilidade_nome):
        """Exporta relatório geral para PDF"""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=18)
        
        story = []
        
        # Título
        story.append(Paragraph(f"Relatório Geral - {contabilidade_nome}", self.styles['CustomTitle']))
        story.append(Spacer(1, 20))
        
        # Data de geração
        story.append(Paragraph(f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}", self.styles['CustomNormal']))
        story.append(Spacer(1, 20))
        
        # Seções do relatório
        for secao, dados in dados_gerais.items():
            story.append(Paragraph(secao.replace('_', ' ').title(), self.styles['CustomHeading']))
            
            if isinstance(dados, list) and dados:
                # Criar tabela para dados em lista
                table_data = []
                if dados:
                    # Usar as chaves do primeiro item como cabeçalho
                    headers = list(dados[0].keys())
                    table_data.append(headers)
                    
                    for item in dados[:20]:  # Limitar a 20 itens por seção
                        row = [str(item.get(header, '')) for header in headers]
                        table_data.append(row)
                
                if table_data:
                    table = Table(table_data, colWidths=[2*inch] * len(headers))
                    table.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('FONTSIZE', (0, 0), (-1, -1), 8),
                        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                        ('GRID', (0, 0), (-1, -1), 1, colors.black)
                    ]))
                    story.append(table)
            
            story.append(Spacer(1, 20))
        
        doc.build(story)
        buffer.seek(0)
        return buffer


class ExcelExporter:
    """Classe para exportação de dados para Excel"""
    
    def export_carteira(self, carteira_data, contabilidade_nome):
        """Exporta dados da carteira para Excel"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Carteira"
        
        # Título
        ws['A1'] = f"Relatório de Carteira - {contabilidade_nome}"
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:D1')
        
        # Data de geração
        ws['A2'] = f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws['A2'].font = Font(size=10)
        
        row = 4
        
        # Resumo geral
        if 'resumo' in carteira_data:
            ws[f'A{row}'] = "Resumo Geral"
            ws[f'A{row}'].font = Font(size=14, bold=True)
            row += 1
            
            resumo = carteira_data['resumo']
            resumo_data = [
                ['Total de Empresas', resumo.get('total_empresas', 0)],
                ['Empresas Ativas', resumo.get('empresas_ativas', 0)],
                ['Empresas Inativas', resumo.get('empresas_inativas', 0)],
                ['Total de Funcionários', resumo.get('total_funcionarios', 0)],
                ['Faturamento Total', resumo.get('faturamento_total', 0)]
            ]
            
            for item in resumo_data:
                ws[f'A{row}'] = item[0]
                ws[f'B{row}'] = item[1]
                row += 1
            
            row += 1
        
        # Lista de empresas
        if 'empresas' in carteira_data:
            ws[f'A{row}'] = "Lista de Empresas"
            ws[f'A{row}'].font = Font(size=14, bold=True)
            row += 1
            
            # Cabeçalhos
            headers = ['Razão Social', 'CNPJ', 'Status', 'Funcionários']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            row += 1
            
            # Dados das empresas
            for empresa in carteira_data['empresas']:
                ws[f'A{row}'] = empresa.get('razao_social', '')
                ws[f'B{row}'] = empresa.get('cnpj', '')
                ws[f'C{row}'] = 'Ativa' if empresa.get('ativo', False) else 'Inativa'
                ws[f'D{row}'] = empresa.get('total_funcionarios', 0)
                row += 1
        
        # Ajustar largura das colunas
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
    
    def export_clientes(self, clientes_data, contabilidade_nome):
        """Exporta dados de clientes para Excel"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Clientes"
        
        # Título
        ws['A1'] = f"Relatório de Clientes - {contabilidade_nome}"
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:D1')
        
        # Data de geração
        ws['A2'] = f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws['A2'].font = Font(size=10)
        
        row = 4
        
        # Cabeçalhos
        headers = ['Nome', 'Documento', 'Faturamento', 'Notas']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        row += 1
        
        # Dados dos clientes
        for cliente in clientes_data:
            ws[f'A{row}'] = cliente.get('nome', '')
            ws[f'B{row}'] = cliente.get('documento', '')
            ws[f'C{row}'] = cliente.get('total_transacoes', 0)
            ws[f'D{row}'] = cliente.get('quantidade_transacoes', 0)
            row += 1
        
        # Ajustar largura das colunas
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
    
    def export_relatorio_geral(self, dados_gerais, contabilidade_nome):
        """Exporta relatório geral para Excel"""
        wb = openpyxl.Workbook()
        
        # Remover planilha padrão
        wb.remove(wb.active)
        
        for secao, dados in dados_gerais.items():
            ws = wb.create_sheet(title=secao.replace('_', ' ').title()[:31])  # Excel limita a 31 caracteres
            
            # Título da seção
            ws['A1'] = f"{secao.replace('_', ' ').title()} - {contabilidade_nome}"
            ws['A1'].font = Font(size=14, bold=True)
            
            if isinstance(dados, list) and dados:
                # Cabeçalhos
                headers = list(dados[0].keys())
                for col, header in enumerate(headers, 1):
                    cell = ws.cell(row=2, column=col, value=header)
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
                
                # Dados
                for row_idx, item in enumerate(dados, 3):
                    for col_idx, header in enumerate(headers, 1):
                        ws.cell(row=row_idx, column=col_idx, value=item.get(header, ''))
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
